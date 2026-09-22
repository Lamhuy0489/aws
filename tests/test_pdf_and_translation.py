import io
import fitz
import pytest
from src.backend.parsers.hybrid_engine import DocumentResult, ParsedPage
from src.backend.exporters.pdf_exporter import PdfExporter
from src.backend.exporters.docx_exporter import DocxExporter
from src.backend.database.db import (
    get_user_settings,
    update_user_settings,
    update_user_password,
    get_user_by_username
)
from src.frontend.server import app

def test_pdf_exporter_generates_valid_pdf():
    """Kiem tra PdfExporter tao tep PDF chuan A4 doc duoc boi PyMuPDF."""
    sample_md = """
# Bien Ban Nghiem Thu Cong Trinh

Day la doan van ban mau tieng Viet co **chu in dam** va *chu in nghieng*.

| STT | Hang Muc Cong Viec | Don Vi | So Luong | Don Gia (VND) |
| :---: | :--- | :---: | :---: | ---: |
| 1 | Kiem thu he thong mang | Goi | 1 | 15,000,000 |
| 2 | Kiem toan ma nguon an toan | Goi | 1 | 25,000,000 |

- Hang muc da duoc phe duyet boi Chu dau tu.
- Cong trinh hoan thanh dung tien do cam ket.
"""
    result = DocumentResult(
        document_id="test-doc-1",
        filename="bien_ban_nghiem_thu.pdf",
        total_pages=1,
        digital_pages_count=1,
        scanned_pages_count=0,
        full_markdown=sample_md,
        processing_time_seconds=1.2,
        pages=[]
    )

    pdf_bytes = PdfExporter.export_to_bytes(result)
    assert pdf_bytes is not None
    assert len(pdf_bytes) > 500

    # Mo lai bang PyMuPDF de xac thuc tinh hop le
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    assert len(doc) >= 1
    page_text = doc[0].get_text()
    assert "Bien Ban Nghiem Thu Cong Trinh" in page_text
    assert "Kiem thu he thong mang" in page_text
    doc.close()

def test_docx_exporter_professional_formatting():
    """Kiem tra DocxExporter khong con de lo ky tu markdown tho va tao bang dung the thuc."""
    sample_md = """
# Thong Bao Phat Hanh Hoa Don

**Nguoi ky:** Nguyen Van A
*Chuc vu:* Ke toan truong

| STT | Dich Vu | Chi Phi |
| :---: | :--- | ---: |
| 1 | May chu EC2 | 500,000 |
"""
    result = DocumentResult(
        document_id="test-doc-2",
        filename="thong_bao.pdf",
        total_pages=1,
        digital_pages_count=1,
        scanned_pages_count=0,
        full_markdown=sample_md,
        processing_time_seconds=0.8,
        pages=[]
    )

    docx_bytes = DocxExporter.export_to_bytes(result)
    assert docx_bytes is not None
    assert len(docx_bytes) > 1000

    # Doc lai tep docx de kiem tra
    from docx import Document
    doc = Document(io.BytesIO(docx_bytes))
    paragraphs_text = [p.text for p in doc.paragraphs]
    full_text = " ".join(paragraphs_text)

    # Khong con de lo ky tu markdown tho nhu ** hoac *
    assert "**Nguoi ky:**" not in full_text
    assert "Nguoi ky:" in full_text

    # Bang bieu duoc tao dung
    assert len(doc.tables) >= 1
    assert len(doc.tables[0].rows) == 2

def test_user_settings_crud():
    """Kiem tra doc va cap nhat cau hinh ca nhan hoa tai khoan."""
    user = get_user_by_username("demo")
    assert user is not None

    # Lay mac dinh
    settings = get_user_settings(user["id"])
    assert settings["theme"] in ["light", "dark"]
    assert settings["language"] in ["vi", "en"]

    # Cap nhat sang dark mode
    updated = update_user_settings(user["id"], {
        "theme": "dark",
        "language": "en",
        "default_model": "gemini-3.6-flash",
        "default_target_lang": "vi",
        "fast_path_default": False
    })
    assert updated["theme"] == "dark"
    assert updated["language"] == "en"
    assert updated["fast_path_default"] == 0

    # Khoi phuc lai ve light mode
    update_user_settings(user["id"], {
        "theme": "light",
        "language": "vi",
        "default_model": "auto",
        "default_target_lang": "vi",
        "fast_path_default": True
    })

def test_api_download_pdf_endpoint():
    """Kiem tra endpoint /api/download/pdf tra ve tep PDF hop le."""
    client = app.test_client()
    user = get_user_by_username("demo")

    with client.session_transaction() as sess:
        sess["user_id"] = user["id"]

    resp = client.post(
        "/api/download/pdf",
        json={
            "document_id": "test-pdf-1",
            "filename": "sample_test.pdf",
            "total_pages": 1,
            "digital_pages_count": 1,
            "scanned_pages_count": 0,
            "full_markdown": "# Test PDF Download\n\nNoi dung xuat PDF kiem thu."
        }
    )

    assert resp.status_code == 200
    assert resp.content_type == "application/pdf"
    assert len(resp.data) > 300
