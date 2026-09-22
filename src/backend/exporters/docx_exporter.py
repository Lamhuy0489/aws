import io
import os
import re
import logging
from typing import List, Tuple
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from src.backend.parsers.hybrid_engine import DocumentResult

logger = logging.getLogger(__name__)

class DocxExporter:
    """Chuyển đổi kết quả bóc tách Markdown thành tài liệu Microsoft Word (.docx) chuẩn thể thức."""

    @staticmethod
    def export_to_bytes(result: DocumentResult) -> bytes:
        """Tạo tài liệu Word chuẩn thể thức từ DocumentResult và trả về mảng bytes."""
        doc = Document()

        # 1. Cấu hình lề trang A4 chuẩn theo Nghị định 30/2020/NĐ-CP
        for section in doc.sections:
            section.page_width = Mm(210)
            section.page_height = Mm(297)
            section.top_margin = Mm(20)
            section.bottom_margin = Mm(20)
            section.left_margin = Mm(30)
            section.right_margin = Mm(20)

        # Cấu hình font chữ mặc định cho tài liệu
        style_normal = doc.styles["Normal"]
        style_normal.font.name = "Times New Roman"
        style_normal.font.size = Pt(13)
        style_normal.font.color.rgb = RGBColor(15, 23, 42)
        style_normal.paragraph_format.line_spacing = 1.2
        style_normal.paragraph_format.space_after = Pt(4)

        # 2. Xử lý từng dòng nội dung Markdown
        lines = result.full_markdown.split("\n")
        table_buffer: List[str] = []
        in_table = False

        for line in lines:
            trimmed = line.strip()

            # Phát hiện khối bảng Markdown
            if trimmed.startswith("|") and trimmed.endswith("|"):
                in_table = True
                table_buffer.append(trimmed)
                continue
            else:
                if in_table and table_buffer:
                    DocxExporter._render_markdown_table_to_docx(doc, table_buffer)
                    table_buffer = []
                    in_table = False

            if not trimmed:
                continue

            # Bỏ qua hoặc xử lý ngắt trang từ comment <!-- Trang X -->
            if trimmed.startswith("<!--") and trimmed.endswith("-->"):
                continue

            # Tiêu đề phân cấp
            if trimmed.startswith("# "):
                h = doc.add_heading(level=1)
                h.paragraph_format.space_before = Pt(12)
                h.paragraph_format.space_after = Pt(6)
                DocxExporter._add_formatted_runs(h, trimmed[2:].strip(), font_size=15, bold=True, color=RGBColor(15, 23, 42))
            elif trimmed.startswith("## "):
                h = doc.add_heading(level=2)
                h.paragraph_format.space_before = Pt(10)
                h.paragraph_format.space_after = Pt(4)
                DocxExporter._add_formatted_runs(h, trimmed[3:].strip(), font_size=13.5, bold=True, color=RGBColor(30, 41, 59))
            elif trimmed.startswith("### "):
                h = doc.add_heading(level=3)
                h.paragraph_format.space_before = Pt(8)
                h.paragraph_format.space_after = Pt(3)
                DocxExporter._add_formatted_runs(h, trimmed[4:].strip(), font_size=12.5, bold=True, italic=True, color=RGBColor(51, 65, 85))
            elif trimmed.startswith("#### "):
                h = doc.add_heading(level=4)
                h.paragraph_format.space_before = Pt(6)
                h.paragraph_format.space_after = Pt(2)
                DocxExporter._add_formatted_runs(h, trimmed[5:].strip(), font_size=12, bold=True, color=RGBColor(71, 85, 105))
            elif trimmed.startswith("> *") and trimmed.endswith("*"):
                # Header trích dẫn đầu trang
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                quote_text = trimmed[3:-1].strip()
                run = p.add_run(quote_text)
                run.italic = True
                run.font.size = Pt(10.5)
                run.font.color.rgb = RGBColor(100, 116, 139)
            elif trimmed.startswith("- ") or trimmed.startswith("* "):
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                DocxExporter._add_formatted_runs(p, trimmed[2:].strip(), font_size=12.5)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.line_spacing = 1.2
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                DocxExporter._add_formatted_runs(p, trimmed, font_size=12.5)

        # Nếu tệp kết thúc bằng một bảng
        if in_table and table_buffer:
            DocxExporter._render_markdown_table_to_docx(doc, table_buffer)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def _add_formatted_runs(paragraph, text: str, font_size: float = 12.5, bold: bool = False, italic: bool = False, color: RGBColor = None):
        """Phân tích các thẻ inline Markdown (**in đậm**, *in nghiêng*, `code`) thành các Run Word định dạng chuẩn."""
        tokens = DocxExporter._tokenize_markdown(text)
        for token_text, is_bold, is_italic, is_code in tokens:
            run = paragraph.add_run(token_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(font_size)
            run.bold = bold or is_bold
            run.italic = italic or is_italic
            if color:
                run.font.color.rgb = color
            if is_code:
                run.font.name = "Consolas"
                run.font.size = Pt(font_size - 1)
                run.font.color.rgb = RGBColor(15, 23, 42)

    @staticmethod
    def _tokenize_markdown(text: str) -> List[Tuple[str, bool, bool, bool]]:
        """Chia văn bản thành các đoạn định dạng: (nội dung, in_đậm, in_nghiêng, code)."""
        pattern = re.compile(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)")
        parts = pattern.split(text)
        tokens = []

        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**") and len(part) >= 4:
                tokens.append((part[2:-2], True, False, False))
            elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
                tokens.append((part[1:-1], False, True, False))
            elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
                tokens.append((part[1:-1], False, False, True))
            else:
                tokens.append((part, False, False, False))

        return tokens

    @staticmethod
    def _render_markdown_table_to_docx(doc: Document, table_lines: List[str]):
        """Dựng bảng Word chuẩn thể thức: Tự co giãn, viền mỏng, shading tiêu đề, lặp header, căn lề cột."""
        rows = []
        col_alignments = []

        for line in table_lines:
            cells = [c.strip() for c in line.strip("|").split("|")]
            # Kiểm tra dòng căn lề :---, :---:, ---:
            if all(set(c).issubset({"-", ":", " "}) for c in cells):
                for c in cells:
                    c_clean = c.strip()
                    if c_clean.startswith(":") and c_clean.endswith(":"):
                        col_alignments.append(WD_ALIGN_PARAGRAPH.CENTER)
                    elif c_clean.endswith(":"):
                        col_alignments.append(WD_ALIGN_PARAGRAPH.RIGHT)
                    else:
                        col_alignments.append(WD_ALIGN_PARAGRAPH.LEFT)
                continue
            rows.append(cells)

        if not rows:
            return

        num_cols = max(len(r) for r in rows)
        if not col_alignments or len(col_alignments) < num_cols:
            col_alignments = [WD_ALIGN_PARAGRAPH.LEFT] * num_cols

        table = doc.add_table(rows=len(rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        # Áp dụng viền kẻ viền mỏng chuyên nghiệp (#CBD5E1)
        DocxExporter._set_table_borders(table)

        for row_idx, row_data in enumerate(rows):
            row = table.rows[row_idx]
            # Ngăn cắt đôi hàng khi rơi vào mép trang (cantSplit)
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(parse_xml(r'<w:cantSplit {}/>'.format(nsdecls('w'))))

            # Hàng tiêu đề: lặp lại khi qua trang mới (tblHeader)
            if row_idx == 0:
                tr_pr.append(parse_xml(r'<w:tblHeader {}/>'.format(nsdecls('w'))))

            for col_idx, cell_value in enumerate(row_data):
                if col_idx < num_cols:
                    cell = row.cells[col_idx]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15

                    align = col_alignments[col_idx] if col_idx < len(col_alignments) else WD_ALIGN_PARAGRAPH.LEFT
                    # Nếu là hàng tiêu đề thì căn giữa
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 else align

                    is_header = (row_idx == 0)
                    DocxExporter._add_formatted_runs(p, cell_value, font_size=11, bold=is_header)

                    # Đổ màu nền cho hàng tiêu đề
                    if row_idx == 0:
                        shading_xml = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
                        cell._tc.get_or_add_tcPr().append(shading_xml)

        doc.add_paragraph()  # Thêm khoảng trống sau bảng

    @staticmethod
    def _set_table_borders(table):
        """Thiết lập viền kẻ mỏng sắc nét (#CBD5E1) cho bảng Word."""
        tbl_pr = table._tbl.tblPr
        borders_xml = parse_xml(
            r'''<w:tblBorders {} >
                <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            </w:tblBorders>'''.format(nsdecls('w'))
        )
        tbl_pr.append(borders_xml)

    @staticmethod
    def export_to_file(result: DocumentResult, output_path: str) -> str:
        """Xuất tệp Word (.docx) ra đường dẫn tệp cụ thể."""
        docx_bytes = DocxExporter.export_to_bytes(result)
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        with open(output_path, "wb") as f:
            f.write(docx_bytes)
        logger.info(f"Đã xuất tệp Word (.docx) chuẩn thể thức tại: {output_path}")
        return output_path
